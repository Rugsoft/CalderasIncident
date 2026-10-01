#!/usr/bin/env python3
"""
generar_word.py
Generador automático del documento de Word (.docx) formal para Calderas CESI
"""

import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def create_document():
    doc = Document()

    # Configuración de márgenes
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Colores corporativos CESI
    COLOR_PRIMARY = RGBColor(13, 34, 56)      # Azul noche
    COLOR_SECONDARY = RGBColor(21, 67, 96)    # Azul corporativo
    COLOR_FLAME = RGBColor(211, 84, 0)        # Naranja fuego
    COLOR_MUTED = RGBColor(86, 101, 115)      # Gris técnico

    # TÍTULO DEL DOCUMENTO
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(4)
    run_title = p_title.add_run("CALDERAS CESI S.L.")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_PRIMARY

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(14)
    run_sub = p_sub.add_run("MEMORIA TÉCNICA DE AMPLIACIONES FUTURAS Y BACKLOG EVOLUTIVO")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(13)
    run_sub.font.bold = True
    run_sub.font.color.rgb = COLOR_FLAME

    # Metadatos del documento
    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_after = Pt(20)
    run_meta = p_meta.add_run("Proyecto: Gestor de Incidencias y Mantenimiento Técnico | Fase: Entrega MVP v1.0 | Fecha: Octubre 2026")
    run_meta.font.name = "Arial"
    run_meta.font.size = Pt(9.5)
    run_meta.font.italic = True
    run_meta.font.color.rgb = COLOR_MUTED

    # Separador visual
    p_line = doc.add_paragraph()
    p_line.paragraph_format.space_after = Pt(16)
    p_line_run = p_line.add_run("_________________________________________________________________________________")
    p_line_run.font.color.rgb = RGBColor(213, 219, 219)

    # 1. INTRODUCCIÓN Y JUSTIFICACIÓN
    h1 = doc.add_heading(level=1)
    run_h1 = h1.add_run("1. Introducción y Justificación del Alcance del MVP")
    run_h1.font.name = "Arial"
    run_h1.font.color.rgb = COLOR_PRIMARY

    p1 = doc.add_paragraph(
        "Para responder con máxima solvencia a los requerimientos de la empresa Calderas CESI S.L., "
        "se ha priorizado la construcción e implantación de un Producto Mínimo Viable (MVP) enfocado en la estabilidad, "
        "la robustez relacional y el blindaje de seguridad. La solución desarrollada ya cubre de forma 100% operativa:"
    )
    p1.paragraph_format.space_after = Pt(6)

    bullets = [
        "Control de Acceso Basado en Roles (RBAC): Perímetros diferenciados para clientes solicitantes, técnicos de mantenimiento y administradores.",
        "Catálogo completo de productos de Calderas CESI: Soporte nativo para calderas murales de condensación (EcoCondens), de alta potencia (TermoMax), de biomasa/pellet (BioPellet) y sistemas híbridos de aerotermia.",
        "Apertura transaccional de incidencias: Manejo atómico ACID con generación correlativa de referencias de servicio (ej. CESI-2026-XXXX).",
        "Blindaje de privacidad y seguridad: Prevención de vulnerabilidades IDOR (Insecure Direct Object Reference), inyecciones SQL con PDO nativo y ataques CSRF con tokens criptográficos.",
        "Pistas inmutables de auditoría: Historial cronológico estricto que registra cada cambio de estado, técnico responsable y motivo de la intervención."
    ]
    for b in bullets:
        bp = doc.add_paragraph(b, style='List Bullet')
        bp.paragraph_format.space_after = Pt(3)

    p_trans = doc.add_paragraph(
        "Aquellas funcionalidades de mayor envergadura contempladas en los manuales de arquitectura técnica "
        "se detallan a continuación para su incorporación planificada en las fases sucesivas de desarrollo."
    )
    p_trans.paragraph_format.space_before = Pt(8)
    p_trans.paragraph_format.space_after = Pt(16)

    # 2. INVENTARIO DE MEJORAS
    h2 = doc.add_heading(level=1)
    run_h2 = h2.add_run("2. Módulos Planificados para Siguientes Fases")
    run_h2.font.name = "Arial"
    run_h2.font.color.rgb = COLOR_PRIMARY

    modulos = [
        (
            "2.1. Custodia Segura de Archivos Adjuntos y Fotografías de Averías (Manual 10)",
            "Permitirá al titular adjuntar fotografías del display de la caldera con códigos de fallo (E01, F28, etc.), "
            "manómetros de presión descompensados o placas de número de serie. A nivel arquitectónico, los binarios se "
            "almacenarán en un directorio privado fuera del DocumentRoot para impedir la ejecución remota de código (RCE), "
            "verificando el tipo MIME real mediante la extensión fileinfo de PHP y distribuyendo las descargas con nombres "
            "hasheados y cabeceras forzadas Content-Disposition: attachment."
        ),
        (
            "2.2. Parte de Trabajo Digital y Firma Biométrica del Cliente en Pantalla",
            "Permitirá al técnico de Calderas CESI, una vez finalizada la reparación o revisión preventiva a domicilio, "
            "cumplimentar en su tablet o smartphone el parte de trabajo oficial. Incluirá la captura de firma biométrica "
            "mediante HTML5 Canvas, el desglose de piezas sustituidas (vasos de expansión, sondas NTC, válvulas de 3 bares) "
            "y la generación automática de un certificado PDF firmado que se enviará de forma inmediata al correo del cliente."
        ),
        (
            "2.3. Sistema de Notificaciones Automáticas (SMS, WhatsApp y Correo)",
            "Para averías clasificadas con prioridad URGENTE (fugas de gas, cortes de calefacción en ola de frío o inundaciones "
            "por rotura de racores), el sistema emitirá avisos instantáneos mediante integración con API de mensajería (Twilio / "
            "WhatsApp Business API) al teléfono del técnico de guardia, reduciendo el tiempo de respuesta a menos de 60 minutos."
        ),
        (
            "2.4. Módulo de Gestión de Stock de Repuestos de Caldera",
            "Conexión directa entre el catálogo de incidencias y el inventario del almacén técnico central de CESI. Cada "
            "intervención descontará automáticamente las piezas empleadas, emitiendo alertas cuando el número de repuestos "
            "críticos caiga por debajo del umbral mínimo de seguridad."
        ),
        (
            "2.5. Automatización de Revisiones Periódicas Obligatorias (RITE)",
            "Implementación de un cron job que monitorice las fechas de instalación de cada caldera registrada en el sistema. "
            "Al cumplirse los 11 meses de la última inspección, el aplicativo generará de forma proactiva una propuesta de cita "
            "para la revisión anual obligatoria, garantizando el cumplimiento normativo de eficiencia energética."
        )
    ]

    for titulo, desc in modulos:
        sub_h = doc.add_heading(level=2)
        sub_run = sub_h.add_run(titulo)
        sub_run.font.name = "Arial"
        sub_run.font.size = Pt(11.5)
        sub_run.font.color.rgb = COLOR_SECONDARY
        sub_h.paragraph_format.space_before = Pt(10)
        sub_h.paragraph_format.space_after = Pt(4)

        mp = doc.add_paragraph(desc)
        mp.paragraph_format.space_after = Pt(10)

    # 3. TABLA DE PRIORIZACIÓN Y CRONOGRAMA
    h3 = doc.add_heading(level=1)
    run_h3 = h3.add_run("3. Matriz de Priorización y Cronograma Estimado")
    run_h3.font.name = "Arial"
    run_h3.font.color.rgb = COLOR_PRIMARY
    h3.paragraph_format.space_before = Pt(14)
    h3.paragraph_format.space_after = Pt(8)

    table = doc.add_table(rows=6, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    headers = ["Módulo / Funcionalidad", "Prioridad", "Estimación", "Fase Programada"]
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        set_cell_background(hdr_cells[i], "154360")
        for p in hdr_cells[i].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(9.5)

    data = [
        ("Subida segura de fotografías/adjuntos", "Alta", "12 horas", "Fase 2.1 (Próxima release)"),
        ("Notificaciones automáticas SMS / Email", "Media-Alta", "8 horas", "Fase 2.2"),
        ("Parte de trabajo digital con firma en tablet", "Media", "16 horas", "Fase 3.1"),
        ("Gestión de stock de repuestos", "Media", "20 horas", "Fase 3.2"),
        ("Alertas de inspección anual RITE", "Estratégica", "14 horas", "Fase 3.3"),
    ]

    for row_idx, row_data in enumerate(data, start=1):
        row_cells = table.rows[row_idx].cells
        bg_color = "F8F9F9" if row_idx % 2 == 0 else "FFFFFF"
        for col_idx, text in enumerate(row_data):
            row_cells[col_idx].text = text
            set_cell_background(row_cells[col_idx], bg_color)
            for p in row_cells[col_idx].paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
                    if col_idx == 0:
                        run.font.bold = True

    # 4. CONCLUSIÓN
    h4 = doc.add_heading(level=1)
    run_h4 = h4.add_run("4. Conclusión Técnica")
    run_h4.font.name = "Arial"
    run_h4.font.color.rgb = COLOR_PRIMARY
    h4.paragraph_format.space_before = Pt(18)
    h4.paragraph_format.space_after = Pt(6)

    p_fin = doc.add_paragraph(
        "El núcleo arquitectónico entregado en este MVP garantiza que todas las extensiones contempladas "
        "podrán incorporarse de forma modular y desacoplada, preservando la inmutabilidad de los datos históricos "
        "y la seguridad del perímetro web de Calderas CESI S.L."
    )
    p_fin.paragraph_format.space_after = Pt(20)

    # Firma y cierre
    p_firma = doc.add_paragraph()
    p_firma.paragraph_format.space_before = Pt(30)
    p_firma.add_run("Calderas CESI S.L. — Departamento de Desarrollo e Implantación Web\n")
    run_subf = p_firma.add_run("Documentación Oficial de Entrega de Proyecto")
    run_subf.font.italic = True
    run_subf.font.size = Pt(9.5)
    run_subf.font.color.rgb = COLOR_MUTED

    ruta_salida = os.path.join(os.path.dirname(__file__), "AMPLIACIONES_FUTURAS_BACKLOG.docx")
    doc.save(ruta_salida)
    print(f"Documento Word generado con éxito en: {ruta_salida}")

if __name__ == "__main__":
    create_document()
